"""The simple-table lane must diagnose native content it cannot preserve."""

import hashlib
import json
import subprocess
import sys
import uuid
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula

import modules.extract.unstructured_xlsx_intake_v1.main as xlsx


def _workbook(path, *, formula=False, merge=False, literal=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Budget"
    sheet.append(["Amount", "Value"])
    sheet.append(["First", 2])
    sheet.append(["Total", "=SUM(B2:B2)" if formula or literal else 2])
    if literal:
        sheet["B3"].data_type = "s"
    if merge:
        sheet.merge_cells("A1:B1")
    book.save(path)
    book.close()
    return path


def _set_cached_value(path, value):
    with ZipFile(path) as archive:
        parts = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    cell = root.find(".//{*}c[@r='B3']")
    cell.set("t", "b" if isinstance(value, bool) else "n")
    cell.find("{*}v").text = str(int(value))
    parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def _replace_part(path, name, transform):
    with ZipFile(path) as archive:
        parts = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    parts[name] = transform(parts[name])
    with ZipFile(path, "w") as archive:
        for part, data in parts.items():
            archive.writestr(part, data)


def _main(monkeypatch, source, outdir):
    monkeypatch.setattr(sys, "argv", [
        "xlsx-intake", "--xlsx", str(source), "--outdir", str(outdir),
    ])
    xlsx.main()


def test_unsupported_workbook_stops_before_bundle(tmp_path):
    run_id = f"xlsx-unsupported-{uuid.uuid4().hex[:8]}"
    run = tmp_path / run_id
    result = subprocess.run(
        [sys.executable, "driver.py", "--recipe", "configs/recipes/recipe-xlsx-html-mvp.yaml",
         "--input-xlsx", "testdata/xlsx-merged-formula.xlsx", "--run-id", run_id,
         "--output-dir", str(run)],
        capture_output=True, text=True, timeout=120,
    )
    assert result.returncode != 0
    report_path = run / "01_unstructured_xlsx_intake_v1/xlsx_source_features.json"
    assert str(report_path) in result.stdout + result.stderr
    report = json.loads(report_path.read_text())
    assert report["disposition"] == "blocked_unsupported_features"
    assert [(row["sheet"], row["cell"], row["expression"]) for row in report["formula_cells"]] == [
        ("Budget", "B5", "=SUM(B3:B4)"), ("Budget", "C5", "=SUM(C3:C4)"),
    ]
    assert report["merged_ranges"] == [{"sheet": "Budget", "range": "A1:C1"}]
    assert all(row["cache_observation"] == "none_observed" for row in report["formula_cells"])
    assert not (run / "01_unstructured_xlsx_intake_v1/elements.jsonl").exists()
    assert not (run / "output/html/manifest.json").exists()
    state = json.loads((run / "pipeline_state.json").read_text())
    assert state["stages"]["xlsx_intake"]["status"] == "failed"


@pytest.mark.parametrize("formula,merge", [(True, False), (False, True), (True, True)])
def test_native_feature_inventory_is_source_located(tmp_path, formula, merge):
    source = _workbook(tmp_path / "source.xlsx", formula=formula, merge=merge)
    report = xlsx.inspect_xlsx_features(source)
    assert report["source_file"] == str(source.resolve())
    assert report["source_sha256"] == hashlib.sha256(source.read_bytes()).hexdigest()
    assert report["disposition"] == "blocked_unsupported_features"
    assert bool(report["formula_cells"]) is formula
    assert bool(report["merged_ranges"]) is merge
    if formula:
        row = report["formula_cells"][0]
        assert (row["sheet"], row["cell"], row["expression"]) == ("Budget", "B3", "=SUM(B2:B2)")
        assert row["cached_xml_value"] is None
        assert row["cache_observation"] == "none_observed"
    if merge:
        assert report["merged_ranges"] == [{"sheet": "Budget", "range": "A1:B1"}]


@pytest.mark.parametrize("cached", [0, False])
def test_cached_formula_is_still_unsupported_and_zero_is_observed(tmp_path, cached):
    source = _workbook(tmp_path / "cached.xlsx", formula=True)
    _set_cached_value(source, cached)
    report = xlsx.inspect_xlsx_features(source)
    assert report["disposition"] == "blocked_unsupported_features"
    row = report["formula_cells"][0]
    assert row["cached_xml_value"] == "0"
    assert row["cached_xml_type"] == ("b" if isinstance(cached, bool) else "n")
    assert row["cache_observation"] == "text_observed"
    assert row["expression"] == "=SUM(B2:B2)"


def test_array_formula_cannot_evade_native_feature_gate(tmp_path):
    source = _workbook(tmp_path / "array.xlsx")
    from openpyxl import load_workbook
    book = load_workbook(source)
    book.active["B2"] = ArrayFormula(ref="B2:B3", text="=A2:A3*2")
    book.save(source)
    book.close()
    report = xlsx.inspect_xlsx_features(source)
    assert report["disposition"] == "blocked_unsupported_features"
    row = report["formula_cells"][0]
    assert row["cell"] == "B2"
    assert row["expression"] == "=A2:A3*2"
    assert row["formula_range"] == "B2:B3"


def test_formula_like_literal_text_is_not_classified_as_formula(tmp_path):
    source = _workbook(tmp_path / "literal.xlsx", literal=True)
    report = xlsx.inspect_xlsx_features(source)
    assert report["disposition"] == "native_feature_gate_clear"
    assert report["formula_cells"] == []
    assert report["merged_ranges"] == []


@pytest.mark.parametrize("formula,merge", [(True, False), (False, True)])
def test_unsupported_features_never_reach_partitioner(tmp_path, monkeypatch, formula, merge):
    source = _workbook(tmp_path / "unsupported.xlsx", formula=formula, merge=merge)
    def forbidden(*args, **kwargs):
        pytest.fail("unsupported native features reached Unstructured")
    monkeypatch.setattr(xlsx, "partition_xlsx_with_unstructured", forbidden)
    with pytest.raises(SystemExit, match="unsupported"):
        _main(monkeypatch, source, tmp_path / "out")
    assert not (tmp_path / "out/elements.jsonl").exists()


def test_source_change_during_partition_prevents_accepted_elements(tmp_path, monkeypatch):
    source = _workbook(tmp_path / "mutable.xlsx")
    def mutate(*args, **kwargs):
        source.write_bytes(source.read_bytes() + b"changed")
        return []
    monkeypatch.setattr(xlsx, "partition_xlsx_with_unstructured", mutate)
    with pytest.raises(SystemExit, match="changed"):
        _main(monkeypatch, source, tmp_path / "out")
    report = json.loads((tmp_path / "out/xlsx_source_features.json").read_text())
    assert report["disposition"] == "blocked_source_changed"
    assert not (tmp_path / "out/elements.jsonl").exists()


def test_unavailable_native_inspection_never_falls_back(tmp_path, monkeypatch):
    source = _workbook(tmp_path / "source.xlsx")
    def unavailable(*args, **kwargs):
        raise ValueError("cannot inspect source")
    monkeypatch.setattr(xlsx, "inspect_xlsx_features", unavailable)
    def forbidden(*args, **kwargs):
        pytest.fail("native inspection failure reached Unstructured")
    monkeypatch.setattr(xlsx, "partition_xlsx_with_unstructured", forbidden)
    with pytest.raises(ValueError, match="cannot inspect"):
        _main(monkeypatch, source, tmp_path / "out")
    assert not (tmp_path / "out/elements.jsonl").exists()


def test_sparse_far_corner_formula_inspects_only_stored_cells(tmp_path):
    source = tmp_path / "sparse.xlsx"
    book = Workbook()
    book.active["A1"] = "Source"
    book.active["XFD1048576"] = "=1+1"
    book.save(source)
    book.close()
    assert source.stat().st_size < 10000
    report = xlsx.inspect_xlsx_features(source)
    assert report["stored_cells_scanned"] == 2
    assert report["disposition"] == "blocked_unsupported_features"
    assert report["formula_cells"][0]["cell"] == "XFD1048576"
    assert report["formula_cells"][0]["expression"] == "=1+1"


def test_huge_merge_is_reported_without_materializing_its_cells(tmp_path):
    source = _workbook(tmp_path / "merge.xlsx", merge=True)
    _replace_part(source, "xl/worksheets/sheet1.xml", lambda data: data.replace(b'A1:B1', b'A1:XFD1048576'))
    report = xlsx.inspect_xlsx_features(source)
    assert report["stored_cells_scanned"] == 5
    assert report["merged_ranges"] == [{"sheet": "Budget", "range": "A1:XFD1048576"}]


def test_shared_formula_without_repeated_text_is_still_blocked(tmp_path):
    source = _workbook(tmp_path / "shared.xlsx", formula=True)
    def shared(data):
        root = ET.fromstring(data)
        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        cell = root.find(".//{*}c[@r='B3']")
        formula = cell.find("{*}f")
        formula.text = None
        formula.set("t", "shared")
        formula.set("si", "7")
        assert cell.tag == f"{{{ns}}}c"
        return ET.tostring(root)
    _replace_part(source, "xl/worksheets/sheet1.xml", shared)
    report = xlsx.inspect_xlsx_features(source)
    assert report["disposition"] == "blocked_unsupported_features"
    row = report["formula_cells"][0]
    assert row["cell"] == "B3" and row["expression"] is None
    assert row["formula_xml_attributes"] == {"t": "shared", "si": "7"}


def test_relative_renamed_worksheet_uses_relationship_owner(tmp_path):
    source = _workbook(tmp_path / "renamed.xlsx", formula=True)
    with ZipFile(source) as archive:
        parts = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    parts["xl/worksheets/renamed.xml"] = parts.pop("xl/worksheets/sheet1.xml")
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b'/xl/worksheets/sheet1.xml', b'worksheets/renamed.xml',
    )
    with ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    report = xlsx.inspect_xlsx_features(source)
    assert report["formula_cells"][0]["sheet"] == "Budget"
    assert report["formula_cells"][0]["cell"] == "B3"


@pytest.mark.parametrize("part", ["xl/workbook.xml", "xl/worksheets/sheet1.xml"])
def test_xml_dtd_is_rejected_before_any_entity_expansion(tmp_path, part):
    source = _workbook(tmp_path / "dtd.xlsx")
    _replace_part(source, part, lambda data: b'<!DOCTYPE source [<!ENTITY value "untrusted">]>' + data)
    with pytest.raises(ValueError, match="DTDs"):
        xlsx.inspect_xlsx_features(source)


def test_external_worksheet_relationship_is_never_followed(tmp_path):
    source = _workbook(tmp_path / "external.xlsx")
    def external(data):
        root = ET.fromstring(data)
        rel = next(r for r in root if r.get("Type", "").endswith("/worksheet"))
        rel.set("TargetMode", "External")
        rel.set("Target", "https://example.invalid/sheet.xml")
        return ET.tostring(root)
    _replace_part(source, "xl/_rels/workbook.xml.rels", external)
    with pytest.raises(ValueError, match="worksheet.*available"):
        xlsx.inspect_xlsx_features(source)


def test_encoded_target_is_inspected_as_same_literal_part_as_consumer(tmp_path):
    from openpyxl import load_workbook
    source = _workbook(tmp_path / "encoded.xlsx", formula=True)
    clean = _workbook(tmp_path / "clean.xlsx")
    with ZipFile(clean) as archive:
        clean_sheet = archive.read("xl/worksheets/sheet1.xml")
    with ZipFile(source) as archive:
        parts = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    parts["xl/worksheets/sheet%31.xml"] = parts["xl/worksheets/sheet1.xml"]
    parts["xl/worksheets/sheet1.xml"] = clean_sheet
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b'/xl/worksheets/sheet1.xml', b'/xl/worksheets/sheet%31.xml',
    )
    with ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    consumer = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    assert consumer["Budget"]["B3"].value == "=SUM(B2:B2)"
    consumer.close()
    report = xlsx.inspect_xlsx_features(source)
    assert report["source_parts"] == [{"sheet": "Budget", "part": "xl/worksheets/sheet%31.xml"}]
    assert report["disposition"] == "blocked_unsupported_features"
    assert report["formula_cells"][0]["cell"] == "B3"


def test_content_type_workbook_authority_matches_consumer(tmp_path):
    from openpyxl import load_workbook
    source = _workbook(tmp_path / "conflicting.xlsx", formula=True)
    clean = _workbook(tmp_path / "clean.xlsx")
    with ZipFile(clean) as archive:
        clean_sheet = archive.read("xl/worksheets/sheet1.xml")
    with ZipFile(source) as archive:
        parts = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    parts["_rels/.rels"] = parts["_rels/.rels"].replace(b'xl/workbook.xml', b'xl/cleanbook.xml')
    parts["xl/cleanbook.xml"] = parts["xl/workbook.xml"]
    parts["xl/_rels/cleanbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b'/xl/worksheets/sheet1.xml', b'/xl/worksheets/clean.xml',
    )
    parts["xl/worksheets/clean.xml"] = clean_sheet
    with ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    consumer = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    assert consumer["Budget"]["B3"].value == "=SUM(B2:B2)"
    consumer.close()
    report = xlsx.inspect_xlsx_features(source)
    assert report["source_parts"] == [{"sheet": "Budget", "part": "xl/worksheets/sheet1.xml"}]
    assert report["disposition"] == "blocked_unsupported_features"
    assert report["formula_cells"][0]["expression"] == "=SUM(B2:B2)"


def test_unavailable_native_part_identity_has_no_guessed_fallback(tmp_path, monkeypatch):
    import openpyxl
    source = _workbook(tmp_path / "no-identity.xlsx")
    load = openpyxl.load_workbook
    def without_identity(*args, **kwargs):
        book = load(*args, **kwargs)
        for sheet in book.worksheets:
            del sheet._worksheet_path
        return book
    monkeypatch.setattr(openpyxl, "load_workbook", without_identity)
    with pytest.raises(ValueError, match="part identity unavailable"):
        xlsx.inspect_xlsx_features(source)
